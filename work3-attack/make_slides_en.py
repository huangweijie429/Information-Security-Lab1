"""Build the English Work 3 presentation (pptx) around the digraph."""
import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = pathlib.Path(__file__).parent
OUT = HERE / "Work3-Presentation-EN.pptx"
NAVY = RGBColor(0x12, 0x39, 0x5C)
GREY = RGBColor(0x52, 0x60, 0x6D)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def title_slide(title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(40)
    s.shapes.title.text_frame.paragraphs[0].font.color.rgb = NAVY
    s.placeholders[1].text = subtitle
    s.placeholders[1].text_frame.paragraphs[0].font.size = Pt(18)
    s.placeholders[1].text_frame.paragraphs[0].font.color.rgb = GREY


def bullets(title, items, size=20):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
    s.shapes.title.text_frame.paragraphs[0].font.color.rgb = NAVY
    body = s.placeholders[1].text_frame
    body.word_wrap = True
    for i, item in enumerate(items):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.space_after = Pt(8)


title_slide("SQL Injection Analysis",
            "Work 3 — modelling an attack as a directed graph\nZhang San · IDU 12345678 · variant 1")

bullets("Topic: SQL injection", [
    "CWE-89; OWASP Top 10 2021 A03: Injection",
    "Input is executed as SQL code -> auth bypass, read/write any table",
    "Precondition: an input point concatenated into SQL; the more DB privileges, the worse",
    "Why it: connects directly to Work 1 and Work 2, and the state machine is clean",
])

bullets("Attacker's sequence of actions", [
    "1  Recon: crawl, port scan, find endpoints that take strings",
    "2  Probe: submit a quote, trigger an SQL error, confirm the injection point",
    "3  Exploit: ' OR 1=1-- to bypass login; UNION SELECT to read the user table",
    "4  Exfiltrate: sqlmap dump, or page the data out",
    "5  Destroy: stacked UPDATE / DROP, change prices or delete tables",
    "6  Blind: boolean/time-based blind injection with SLEEP(), exhausting the pool",
])

s = prs.slides.add_slide(prs.slide_layouts[6])
s.shapes.add_picture(str(HERE / "digraph-en.png"), Inches(0.35), Inches(0.45), width=Inches(12.6))
tx = s.shapes.add_textbox(Inches(0.35), Inches(0.05), Inches(12.6), Inches(0.4))
tx.text_frame.text = "Directed graph: S0 -> Si -> Sk (each edge = one attacker action)"
tx.text_frame.paragraphs[0].font.size = Pt(20)
tx.text_frame.paragraphs[0].font.color.rgb = NAVY

bullets("States and evidence", [
    "S0 normal -> S1 recon -> S2 injection point found -> S3 injection succeeded",
    "S3 -> S4 dump data -> Sk1 information theft",
    "S3 -> S5 tamper/delete -> Sk2 information modified or lost",
    "S2 -> S6 blind/heavy queries -> Sk3 denial of service",
    "Every state has observable evidence: logs, errors, slow queries, traffic",
], size=19)

bullets("Countermeasures: cut the edges", [
    "Prevent: parameterized queries, allow-lists, least-privilege DB account, no verbose errors",
    "Detect: WAF rules, database auditing, slow-query and connection monitoring",
    "Respond: force password resets and enable 2FA, restore from backup, rate limit",
    "Priority: edges closest to S0 are the cheapest to harden with the highest payoff",
])

bullets("Conclusion", [
    "SQL injection is a broken boundary between data and code",
    "A directed graph turns a vague narrative into verifiable states and actions",
    "Every edge is both an attack step and a place to break the attack",
    "The same attack can be drawn differently, but S0 and Sk must stay stable",
], size=22)

prs.save(str(OUT))
print("saved", OUT)
