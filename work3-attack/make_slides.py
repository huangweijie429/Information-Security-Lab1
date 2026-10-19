"""Build the Work 3 presentation (pptx) around the digraph."""
import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

HERE = pathlib.Path(__file__).parent
OUT = HERE / "Work3-Attack-Presentation.pptx"
NAVY = RGBColor(0x12, 0x39, 0x5C)
GREY = RGBColor(0x52, 0x60, 0x6D)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def title_slide(title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(40)
    s.shapes.title.text_frame.paragraphs[0].font.color.rgb = NAVY
    s.placeholders[1].text = subtitle
    s.placeholders[1].text_frame.paragraphs[0].font.size = Pt(18)
    s.placeholders[1].text_frame.paragraphs[0].font.color.rgb = GREY


def bullets(title, items, size=20):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
    s.shapes.title.text_frame.paragraphs[0].font.color.rgb = NAVY
    body = s.placeholders[1].text_frame
    body.word_wrap = True
    for i, item in enumerate(items):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.space_after = Pt(8)


title_slide("SQL 注入攻击分析",
            "Work 3 — 计算机攻击的有向图建模\nZhang San · IDU 12345678 · variant 1")

bullets("选题：SQL 注入", [
    "CWE-89；OWASP Top 10 2021 A03: Injection",
    "输入被当作 SQL 代码执行 → 绕过认证、读写任意表",
    "前提：存在拼接 SQL 的输入点；数据库账号权限越大后果越重",
    "为什么选它：与 Work 1/2 的知识点直接衔接，状态机清晰",
])

bullets("攻击者动作序列", [
    "1  踩点：爬站、端口扫描，找出接受字符串参数的接口",
    "2  探测：提交单引号，触发 SQL 报错，确认注入点",
    "3  利用：' OR 1=1-- 绕过登录；UNION SELECT 读取用户表",
    "4  外泄：sqlmap 批量 dump，或分页带出数据",
    "5  破坏：堆叠查询 UPDATE / DROP，篡改价格或删表",
    "6  盲注：布尔/时间盲注 + SLEEP()，同时拖垮连接池",
])

# full-width digraph slide
s = prs.slides.add_slide(prs.slide_layouts[6])
s.shapes.add_picture(str(HERE / "digraph.png"), Inches(0.35), Inches(0.45), width=Inches(12.6))
tx = s.shapes.add_textbox(Inches(0.35), Inches(0.05), Inches(12.6), Inches(0.4))
tx.text_frame.text = "有向图：S0 → Si → Sk（每条边 = 一个攻击者动作）"
tx.text_frame.paragraphs[0].font.size = Pt(20)
tx.text_frame.paragraphs[0].font.color.rgb = NAVY

bullets("状态与判定", [
    "S0 正常 → S1 侦察 → S2 确认注入点 → S3 注入成功",
    "S3 → S4 导出数据 → Sk1 信息窃取",
    "S3 → S5 篡改/删除 → Sk2 信息被修改或丢失",
    "S2 → S6 盲注/重查询 → Sk3 拒绝服务",
    "每个状态都有可观测判据：日志、错误、慢查询、流量",
], size=19)

bullets("对策：按边切断", [
    "预防：参数化查询、白名单、最小权限数据库账号、关闭详细报错",
    "检测：WAF 规则、数据库审计、慢查询与连接数监控",
    "响应：强制重置口令并启用 2FA、从备份恢复、限流熔断",
    "优先级：越靠近 S0 的边，加固成本越低、收益越大",
])

bullets("结论", [
    "SQL 注入的本质是「数据与代码的边界被打破」",
    "有向图把模糊的叙事拆成可验证的状态与动作",
    "每条边既是一个攻击步骤，也是一个可加固的断点",
    "同一攻击可因策略不同画出不同的图，但 S0 与 Sk 必须稳定定义",
], size=22)

prs.save(str(OUT))
print("saved", OUT)
